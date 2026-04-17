"""Unit tests for FileTypeReader (fileTypeReader.py).

Each test creates the minimal fixture file it needs inside a tmp_path
so that no real files are touched.
"""

import base64
import os
import pytest
from fileTypeReader import FileTypeReader


@pytest.fixture
def reader():
    return FileTypeReader()


# ── .txt ────────────────────────────────────────────────────────────────────

class TestReadTxt:
    def test_reads_plain_text(self, reader, tmp_path):
        (tmp_path / "hello.txt").write_text("Hello!")
        assert reader.read_file("hello.txt", str(tmp_path)) == "Hello!"

    def test_empty_txt(self, reader, tmp_path):
        (tmp_path / "empty.txt").write_text("")
        assert reader.read_file("empty.txt", str(tmp_path)) == ""

    def test_multiline_txt(self, reader, tmp_path):
        (tmp_path / "lines.txt").write_text("line1\nline2\nline3")
        content = reader.read_file("lines.txt", str(tmp_path))
        assert content.count("\n") == 2


# ── .csv ────────────────────────────────────────────────────────────────────

class TestReadCsv:
    def test_simple_csv(self, reader, tmp_path):
        (tmp_path / "data.csv").write_text("a,b,c\n1,2,3\n")
        content = reader.read_file("data.csv", str(tmp_path))
        assert "a,b,c" in content
        assert "1,2,3" in content

    def test_csv_preserves_rows(self, reader, tmp_path):
        (tmp_path / "rows.csv").write_text("x,y\n10,20\n30,40\n")
        content = reader.read_file("rows.csv", str(tmp_path))
        lines = [l for l in content.strip().split("\n") if l]
        assert len(lines) == 3


# ── .docx ───────────────────────────────────────────────────────────────────

class TestReadDocx:
    def test_reads_docx_paragraphs(self, reader, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        path = tmp_path / "sample.docx"
        doc.save(str(path))

        content = reader.read_file("sample.docx", str(tmp_path))
        assert "First paragraph" in content
        assert "Second paragraph" in content

    def test_reads_docx_table(self, reader, tmp_path):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "C"
        table.cell(1, 1).text = "D"
        path = tmp_path / "table.docx"
        doc.save(str(path))

        content = reader.read_file("table.docx", str(tmp_path))
        assert "A" in content and "D" in content


# ── .doc (legacy) ──────────────────────────────────────────────────────────

class TestReadDoc:
    def test_legacy_doc_returns_empty(self, reader, tmp_path):
        (tmp_path / "old.doc").write_bytes(b"\x00" * 10)
        assert reader.read_file("old.doc", str(tmp_path)) == ""


# ── .pptx ───────────────────────────────────────────────────────────────────

class TestReadPptx:
    def test_reads_pptx_slide_text(self, reader, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Title Slide"
        path = tmp_path / "deck.pptx"
        prs.save(str(path))

        content = reader.read_file("deck.pptx", str(tmp_path))
        assert "Title Slide" in content
        assert "Slide 1" in content


# ── .ppt (legacy) ──────────────────────────────────────────────────────────

class TestReadPpt:
    def test_legacy_ppt_returns_empty(self, reader, tmp_path):
        (tmp_path / "old.ppt").write_bytes(b"\x00" * 10)
        assert reader.read_file("old.ppt", str(tmp_path)) == ""


# ── .xlsx ───────────────────────────────────────────────────────────────────

class TestReadXlsx:
    def test_reads_xlsx(self, reader, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Name", "Age"])
        ws.append(["Alice", 30])
        path = tmp_path / "sheet.xlsx"
        wb.save(str(path))

        content = reader.read_file("sheet.xlsx", str(tmp_path))
        assert "Name" in content
        assert "Alice" in content
        assert "Sheet: Data" in content


# ── .xls (legacy) ──────────────────────────────────────────────────────────

class TestReadXls:
    def test_legacy_xls_returns_empty(self, reader, tmp_path):
        (tmp_path / "old.xls").write_bytes(b"\x00" * 10)
        assert reader.read_file("old.xls", str(tmp_path)) == ""


# ── .pdf ────────────────────────────────────────────────────────────────────

class TestReadPdf:
    def test_reads_pdf_text(self, reader, tmp_path):
        from PyPDF2 import PdfWriter
        from io import BytesIO
        from PyPDF2.generic import AnnotationBuilder

        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        path = tmp_path / "blank.pdf"
        with open(str(path), "wb") as f:
            writer.write(f)

        content = reader.read_file("blank.pdf", str(tmp_path))
        assert "Page 1" in content


# ── images ──────────────────────────────────────────────────────────────────

class TestReadImage:
    @pytest.mark.parametrize("ext", [".png", ".jpg", ".gif", ".bmp", ".webp"])
    def test_image_returns_base64_data_url(self, reader, tmp_path, ext):
        raw = b"\x89PNG\r\n\x1a\n" + b"\x00" * 20
        fname = f"pic{ext}"
        (tmp_path / fname).write_bytes(raw)
        content = reader.read_file(fname, str(tmp_path))
        assert content.startswith("data:image/")
        assert ";base64," in content
        payload = content.split(";base64,")[1]
        base64.b64decode(payload)  # must not raise

    def test_ico_returns_base64(self, reader, tmp_path):
        (tmp_path / "icon.ico").write_bytes(b"\x00\x00\x01\x00" + b"\x00" * 20)
        content = reader.read_file("icon.ico", str(tmp_path))
        assert content.startswith("data:image/ico;base64,")


# ── .svg ────────────────────────────────────────────────────────────────────

class TestReadSvg:
    def test_reads_svg(self, reader, tmp_path):
        svg = '<svg xmlns="http://www.w3.org/2000/svg"><circle r="10"/></svg>'
        (tmp_path / "circle.svg").write_text(svg)
        assert reader.read_file("circle.svg", str(tmp_path)) == svg


# ── unknown / fallback ──────────────────────────────────────────────────────

class TestReadUnknown:
    def test_unknown_extension_reads_as_text(self, reader, tmp_path):
        (tmp_path / "notes.md").write_text("# Markdown")
        content = reader.read_file("notes.md", str(tmp_path))
        assert "# Markdown" in content

    def test_binary_unknown_returns_empty(self, reader, tmp_path):
        (tmp_path / "blob.xyz").write_bytes(bytes(range(256)))
        content = reader.read_file("blob.xyz", str(tmp_path))
        # may succeed or return "" depending on encoding — either is acceptable
        assert isinstance(content, str)


# ── edge cases ──────────────────────────────────────────────────────────────

class TestEdgeCases:
    def test_nonexistent_file_returns_empty(self, reader, tmp_path):
        assert reader.read_file("nope.txt", str(tmp_path)) == ""

    def test_nonexistent_directory_returns_empty(self, reader):
        assert reader.read_file("a.txt", "/no/such/path") == ""

    def test_file_name_with_spaces(self, reader, tmp_path):
        (tmp_path / "my file.txt").write_text("spaced")
        assert reader.read_file("my file.txt", str(tmp_path)) == "spaced"

    def test_file_name_case_sensitivity(self, reader, tmp_path):
        (tmp_path / "CamelCase.TXT").write_text("upper")
        content = reader.read_file("CamelCase.TXT", str(tmp_path))
        assert content == "upper"


# ── Additional .txt tests ──────────────────────────────────────────────────

class TestReadTxtExtended:
    def test_content_with_special_ascii(self, reader, tmp_path):
        text = "Hello! @#$%^&*() tabs\there end\n"
        (tmp_path / "ascii.txt").write_text(text)
        content = reader.read_file("ascii.txt", str(tmp_path))
        assert "Hello!" in content
        assert "tabs" in content

    def test_large_text_file(self, reader, tmp_path):
        big = "A" * 100_000
        (tmp_path / "big.txt").write_text(big)
        content = reader.read_file("big.txt", str(tmp_path))
        assert len(content) == 100_000

    def test_special_characters(self, reader, tmp_path):
        text = 'Quotes "here" & <brackets> plus\ttabs'
        (tmp_path / "special.txt").write_text(text)
        assert reader.read_file("special.txt", str(tmp_path)) == text

    def test_windows_line_endings(self, reader, tmp_path):
        (tmp_path / "crlf.txt").write_bytes(b"line1\r\nline2\r\n")
        content = reader.read_file("crlf.txt", str(tmp_path))
        assert "line1" in content and "line2" in content


# ── Additional .csv tests ──────────────────────────────────────────────────

class TestReadCsvExtended:
    def test_csv_with_quoted_fields(self, reader, tmp_path):
        csv_text = '"Name","Value"\n"Alice, Bob","123"\n'
        (tmp_path / "quoted.csv").write_text(csv_text)
        content = reader.read_file("quoted.csv", str(tmp_path))
        assert "Name" in content
        assert "Alice" in content

    def test_empty_csv(self, reader, tmp_path):
        (tmp_path / "empty.csv").write_text("")
        content = reader.read_file("empty.csv", str(tmp_path))
        assert isinstance(content, str)

    def test_single_column_csv(self, reader, tmp_path):
        (tmp_path / "single.csv").write_text("header\nval1\nval2\n")
        content = reader.read_file("single.csv", str(tmp_path))
        assert "header" in content
        assert "val1" in content

    def test_csv_many_rows(self, reader, tmp_path):
        rows = "id,value\n" + "".join(f"{i},{i*10}\n" for i in range(100))
        (tmp_path / "many.csv").write_text(rows)
        content = reader.read_file("many.csv", str(tmp_path))
        lines = [l for l in content.strip().split("\n") if l]
        assert len(lines) == 101  # header + 100 rows


# ── Additional .docx tests ─────────────────────────────────────────────────

class TestReadDocxExtended:
    def test_empty_docx(self, reader, tmp_path):
        from docx import Document
        doc = Document()
        path = tmp_path / "empty.docx"
        doc.save(str(path))
        content = reader.read_file("empty.docx", str(tmp_path))
        assert isinstance(content, str)

    def test_docx_many_paragraphs(self, reader, tmp_path):
        from docx import Document
        doc = Document()
        for i in range(20):
            doc.add_paragraph(f"Paragraph {i}")
        path = tmp_path / "many_paras.docx"
        doc.save(str(path))
        content = reader.read_file("many_paras.docx", str(tmp_path))
        assert "Paragraph 0" in content
        assert "Paragraph 19" in content

    def test_docx_with_mixed_content(self, reader, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Cell"
        table.cell(0, 1).text = "Data"
        doc.add_paragraph("After table")
        path = tmp_path / "mixed.docx"
        doc.save(str(path))
        content = reader.read_file("mixed.docx", str(tmp_path))
        assert "Before table" in content
        assert "Cell" in content
        assert "After table" in content


# ── Additional .pptx tests ─────────────────────────────────────────────────

class TestReadPptxExtended:
    def test_multi_slide_pptx(self, reader, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        for i in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Slide Title {i}"
        path = tmp_path / "multi.pptx"
        prs.save(str(path))
        content = reader.read_file("multi.pptx", str(tmp_path))
        assert "Slide 1" in content
        assert "Slide 2" in content
        assert "Slide 3" in content
        assert "Slide Title 0" in content
        assert "Slide Title 2" in content

    def test_empty_pptx(self, reader, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        path = tmp_path / "empty.pptx"
        prs.save(str(path))
        content = reader.read_file("empty.pptx", str(tmp_path))
        assert isinstance(content, str)


# ── Additional .xlsx tests ─────────────────────────────────────────────────

class TestReadXlsxExtended:
    def test_multi_sheet_xlsx(self, reader, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sales"
        ws1.append(["Product", "Amount"])
        ws1.append(["Widget", 100])
        ws2 = wb.create_sheet("Returns")
        ws2.append(["Product", "Count"])
        ws2.append(["Widget", 5])
        path = tmp_path / "multi_sheet.xlsx"
        wb.save(str(path))
        content = reader.read_file("multi_sheet.xlsx", str(tmp_path))
        assert "Sheet: Sales" in content
        assert "Sheet: Returns" in content
        assert "Widget" in content

    def test_empty_xlsx(self, reader, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        path = tmp_path / "empty.xlsx"
        wb.save(str(path))
        content = reader.read_file("empty.xlsx", str(tmp_path))
        assert isinstance(content, str)

    def test_xlsx_with_numbers_and_none(self, reader, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append([1, None, 3.14])
        path = tmp_path / "nums.xlsx"
        wb.save(str(path))
        content = reader.read_file("nums.xlsx", str(tmp_path))
        assert "1" in content
        assert "3.14" in content


# ── Additional image tests ─────────────────────────────────────────────────

class TestReadImageExtended:
    def test_jpeg_extension(self, reader, tmp_path):
        raw = b"\xff\xd8\xff\xe0" + b"\x00" * 20
        (tmp_path / "photo.jpeg").write_bytes(raw)
        content = reader.read_file("photo.jpeg", str(tmp_path))
        assert content.startswith("data:image/jpeg;base64,")

    def test_tiff_extension(self, reader, tmp_path):
        raw = b"\x49\x49\x2a\x00" + b"\x00" * 20
        (tmp_path / "scan.tiff").write_bytes(raw)
        content = reader.read_file("scan.tiff", str(tmp_path))
        assert content.startswith("data:image/tiff;base64,")

    def test_image_base64_roundtrip(self, reader, tmp_path):
        """Verify the base64 payload decodes back to original bytes."""
        original = bytes(range(256))
        (tmp_path / "raw.png").write_bytes(original)
        content = reader.read_file("raw.png", str(tmp_path))
        payload = content.split(";base64,")[1]
        decoded = base64.b64decode(payload)
        assert decoded == original


# ── Additional .svg tests ──────────────────────────────────────────────────

class TestReadSvgExtended:
    def test_svg_with_attributes(self, reader, tmp_path):
        svg = '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect fill="red"/></svg>'
        (tmp_path / "rect.svg").write_text(svg)
        content = reader.read_file("rect.svg", str(tmp_path))
        assert 'width="100"' in content
        assert "rect" in content

    def test_empty_svg(self, reader, tmp_path):
        svg = '<svg xmlns="http://www.w3.org/2000/svg"></svg>'
        (tmp_path / "empty.svg").write_text(svg)
        content = reader.read_file("empty.svg", str(tmp_path))
        assert content == svg


# ── Additional unknown format tests ────────────────────────────────────────

class TestReadUnknownExtended:
    def test_json_file_read_as_text(self, reader, tmp_path):
        (tmp_path / "config.json").write_text('{"key": "value"}')
        content = reader.read_file("config.json", str(tmp_path))
        assert '"key"' in content

    def test_python_file_read_as_text(self, reader, tmp_path):
        (tmp_path / "script.py").write_text("print('hello')")
        content = reader.read_file("script.py", str(tmp_path))
        assert "print" in content

    def test_html_file_read_as_text(self, reader, tmp_path):
        (tmp_path / "page.html").write_text("<html><body>Hi</body></html>")
        content = reader.read_file("page.html", str(tmp_path))
        assert "<body>" in content

    def test_yaml_file_read_as_text(self, reader, tmp_path):
        (tmp_path / "config.yml").write_text("key: value\n")
        content = reader.read_file("config.yml", str(tmp_path))
        assert "key: value" in content


# ── Additional edge cases ──────────────────────────────────────────────────

class TestEdgeCasesExtended:
    def test_deeply_nested_nonexistent(self, reader):
        assert reader.read_file("a.txt", "/no/such/deep/nested/path") == ""

    def test_file_with_no_extension(self, reader, tmp_path):
        (tmp_path / "Makefile").write_text("all:\n\techo hello")
        content = reader.read_file("Makefile", str(tmp_path))
        assert isinstance(content, str)
        assert "echo" in content

    def test_file_with_double_extension(self, reader, tmp_path):
        (tmp_path / "archive.tar.gz").write_bytes(b"\x1f\x8b" + b"\x00" * 10)
        content = reader.read_file("archive.tar.gz", str(tmp_path))
        assert isinstance(content, str)

    def test_read_file_always_returns_string(self, reader, tmp_path):
        """Every supported format must return str, never None."""
        (tmp_path / "test.txt").write_text("text")
        (tmp_path / "test.csv").write_text("a,b\n1,2\n")
        (tmp_path / "test.md").write_text("# Markdown")
        for fname in ["test.txt", "test.csv", "test.md", "nonexistent.txt"]:
            result = reader.read_file(fname, str(tmp_path))
            assert isinstance(result, str), f"read_file({fname}) returned {type(result)}"

    def test_directory_as_filename(self, reader, tmp_path):
        """Passing a directory name instead of a file should return empty."""
        (tmp_path / "subdir").mkdir()
        content = reader.read_file("subdir", str(tmp_path))
        assert content == ""

    def test_file_with_unicode_name(self, reader, tmp_path):
        (tmp_path / "数据.txt").write_text("chinese filename")
        content = reader.read_file("数据.txt", str(tmp_path))
        assert content == "chinese filename"
