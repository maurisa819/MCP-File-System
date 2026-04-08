import logging
import os
import csv
import re
import xml.etree.ElementTree as ET
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook
from PyPDF2 import PdfReader

logger = logging.getLogger("fastmcp.server")


class FileTypeReader:
    """Handles reading various file types and extracting their content."""

    def read_file(self, file_name: str, directory: str) -> str:
        """Reads and returns the content of a file."""
        logger.info(f"Reading file content from {file_name} in directory {directory}")

        file_path = os.path.join(directory, file_name)

        try:
            if not os.path.isfile(file_path):
                logger.error(f"File not found: {file_path}")
                return ""

            _, file_ext = os.path.splitext(file_name)
            file_ext = file_ext.lower()

            if file_ext == ".txt":
                return self._clean_text(self._read_txt(file_path))

            elif file_ext == ".csv":
                return self._clean_text(self._read_csv(file_path))

            elif file_ext == ".docx":
                return self._clean_text(self._read_docx(file_path))

            elif file_ext == ".doc":
                logger.warning("Legacy .doc format is not supported. Please convert it to .docx.")
                return "Legacy .doc format is not supported. Please convert it to .docx."

            elif file_ext == ".pptx":
                return self._clean_text(self._read_pptx(file_path))

            elif file_ext == ".ppt":
                logger.warning("Legacy .ppt format is not supported. Please convert it to .pptx.")
                return "Legacy .ppt format is not supported. Please convert it to .pptx."

            elif file_ext == ".xlsx":
                return self._clean_text(self._read_xlsx(file_path))

            elif file_ext == ".xls":
                logger.warning("Legacy .xls format is not supported. Please convert it to .xlsx.")
                return "Legacy .xls format is not supported. Please convert it to .xlsx."

            elif file_ext == ".pdf":
                return self._clean_text(self._read_pdf(file_path))

            elif file_ext in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tiff", ".ico"]:
                return self._read_image(file_path, file_ext)

            elif file_ext == ".svg":
                return self._clean_text(self._read_svg(file_path))

            else:
                return self._clean_text(self._read_unknown(file_path))

        except Exception as e:
            logger.error(f"Error reading file {file_path}: {e}")
            return ""

    def _clean_text(self, text: str) -> str:
        """Normalizes extracted text for better summarization."""
        if not text:
            return ""

        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"[^\S\n\t]+", " ", text)
        text = re.sub(r"[\x00-\x08\x0B-\x1F\x7F]", "", text)

        lines = [line.strip() for line in text.split("\n")]

        cleaned = []
        empty_count = 0
        for line in lines:
            if not line:
                empty_count += 1
                if empty_count <= 1:
                    cleaned.append("")
            else:
                empty_count = 0
                cleaned.append(line)

        text = "\n".join(cleaned).strip()
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text

    def _read_txt(self, file_path: str) -> str:
        """Reads a text file with encoding fallbacks."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue
            except Exception as e:
                logger.error(f"Error reading TXT file {file_path} with {encoding}: {e}")
                return ""

        logger.error(f"Could not decode TXT file {file_path}")
        return ""

    def _read_csv(self, file_path: str) -> str:
        """Reads a CSV file with delimiter sniffing and encoding fallbacks."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]

        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding, newline="") as file:
                    sample = file.read(4096)
                    file.seek(0)

                    try:
                        dialect = csv.Sniffer().sniff(sample)
                    except Exception:
                        dialect = csv.excel

                    reader = csv.reader(file, dialect)
                    content = []

                    for i, row in enumerate(reader, start=1):
                        row_content = [str(cell).strip() for cell in row]
                        content.append(" | ".join(row_content))

                        # soft limit for extremely large csvs
                        if i >= 500:
                            content.append("[Output truncated after 500 rows]")
                            break

                    return "\n".join(content)

            except UnicodeDecodeError:
                continue
            except Exception as e:
                logger.error(f"Error reading CSV file {file_path} with {encoding}: {e}")
                return ""

        logger.error(f"Could not decode CSV file {file_path}")
        return ""

    def _read_docx(self, file_path: str) -> str:
        """Reads a DOCX file including paragraphs and tables."""
        try:
            doc = DocxDocument(file_path)
            content = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    content.append(text)

            for table_num, table in enumerate(doc.tables, start=1):
                content.append(f"--- Table {table_num} ---")
                for row in table.rows:
                    row_content = [cell.text.strip() for cell in row.cells]
                    content.append(" | ".join(row_content))

            if not content:
                return "[No readable text found in DOCX file]"

            return "\n".join(content)

        except Exception as e:
            logger.error(f"Error reading DOCX file {file_path}: {e}")
            return ""

    def _read_pptx(self, file_path: str) -> str:
        """Reads a PPTX file with better slide structure."""
        try:
            prs = Presentation(file_path)
            content = []

            for slide_num, slide in enumerate(prs.slides, start=1):
                content.append(f"--- Slide {slide_num} ---")
                slide_texts = []

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text = shape.text.strip()
                        if text:
                            slide_texts.append(text)

                if slide_texts:
                    content.extend(slide_texts)
                else:
                    content.append("[No readable text on this slide]")

            return "\n".join(content)

        except Exception as e:
            logger.error(f"Error reading PPTX file {file_path}: {e}")
            return ""

    def _read_xlsx(self, file_path: str) -> str:
        """Reads an XLSX file with sheet names and row limits."""
        try:
            wb = load_workbook(file_path, data_only=True)
            content = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                content.append(f"--- Sheet: {sheet_name} ---")

                row_count = 0
                for row in ws.iter_rows(values_only=True):
                    row_content = [str(cell).strip() if cell is not None else "" for cell in row]
                    if any(cell != "" for cell in row_content):
                        content.append(" | ".join(row_content))
                        row_count += 1

                    if row_count >= 300:
                        content.append("[Output truncated after 300 non-empty rows]")
                        break

            if not content:
                return "[No readable spreadsheet content found]"

            return "\n".join(content)

        except Exception as e:
            logger.error(f"Error reading XLSX file {file_path}: {e}")
            return ""

    def _read_pdf(self, file_path: str) -> str:
        """Reads a PDF file using PyPDF2 text extraction."""
        try:
            pdf_reader = PdfReader(file_path)
            content = []

            for page_num, page in enumerate(pdf_reader.pages, start=1):
                content.append(f"--- Page {page_num} ---")

                try:
                    page_text = page.extract_text()
                except Exception as page_error:
                    logger.warning(f"Could not extract text from page {page_num} in {file_path}: {page_error}")
                    page_text = ""

                if page_text and page_text.strip():
                    content.append(page_text.strip())
                else:
                    content.append("[No readable text extracted from this page]")

            return "\n".join(content)

        except Exception as e:
            logger.error(f"Error reading PDF file {file_path}: {e}")
            return ""

    def _read_image(self, file_path: str, file_ext: str) -> str:
        """Returns a descriptive placeholder instead of raw base64."""
        try:
            file_size = os.path.getsize(file_path)
            return (
                f"[Image file detected: {os.path.basename(file_path)} | "
                f"type={file_ext} | size_bytes={file_size}. "
                "Direct text extraction is not available for this image in the current reader.]"
            )
        except Exception as e:
            logger.error(f"Error reading image file {file_path}: {e}")
            return ""

    def _read_svg(self, file_path: str) -> str:
        """Reads an SVG file and tries to extract human-readable text."""
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                svg_content = file.read()

            text_parts = []

            try:
                root = ET.fromstring(svg_content)
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        text_parts.append(elem.text.strip())
            except Exception:
                # fallback to raw content if XML parsing fails
                return svg_content

            if text_parts:
                return "\n".join(text_parts)

            return svg_content

        except Exception as e:
            logger.error(f"Error reading SVG file {file_path}: {e}")
            return ""

    def _read_unknown(self, file_path: str) -> str:
        """Attempts to read an unknown file format as plain text."""
        logger.warning(f"Unsupported file format: {file_path}")

        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue
            except Exception:
                return ""

        return ""